"""합성 보안 점검 CSV를 검증하고 Excel 보고서로 정리한다."""
from __future__ import annotations

import argparse
from collections import Counter
import csv
from datetime import date
import hashlib
import io
import json
from pathlib import Path
import re


COLUMNS = ("date", "asset", "check_id", "severity", "status", "evidence")
SEVERITIES = ("high", "medium", "low", "info")
STATUSES = ("pass", "fail", "review", "na")


def load_checks(path: Path):
    source = path.read_bytes()
    reader = csv.DictReader(io.StringIO(source.decode("utf-8-sig"), newline=""), strict=True)
    if reader.fieldnames != list(COLUMNS):
        raise ValueError("CSV 헤더와 순서는 " + ",".join(COLUMNS) + "여야 합니다.")
    records, issues, seen = [], [], set()
    for number, raw in enumerate(reader, 2):
        reasons = []
        row = {key: (raw.get(key) or "").strip() for key in COLUMNS}
        if None in raw or any(raw.get(key) is None for key in COLUMNS):
            reasons.append("열 수 불일치")
        try:
            if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", row["date"]):
                raise ValueError
            parsed_date = date.fromisoformat(row["date"])
        except ValueError:
            reasons.append("date는 유효한 YYYY-MM-DD여야 함")
        for key in ("asset", "check_id", "evidence"):
            if not row[key]:
                reasons.append(f"{key}가 비어 있음")
        if row["severity"] not in SEVERITIES:
            reasons.append("severity 허용값 오류")
        if row["status"] not in STATUSES:
            reasons.append("status 허용값 오류")
        # Excel이 저장할 수 없는 제어 문자와 긴 셀을 입력 단계에서 분리한다.
        if any(re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value)
               or len(value) > 4000 for value in row.values()):
            reasons.append("금지 제어 문자 또는 셀 길이 4000자 초과")
        key = (row["date"], row["asset"], row["check_id"])
        if key in seen:
            reasons.append("동일 date/asset/check_id 중복")
        if reasons:
            # 원본 전체는 CSV에 보존한다. JSON escaping으로 제어 문자도 기록한다.
            issues.append((number, "; ".join(reasons), json.dumps(raw, ensure_ascii=True)))
        else:
            seen.add(key)
            records.append({**row, "date": parsed_date, "source_record": number})
    return records, issues, hashlib.sha256(source).hexdigest()


def summarize(records, issues):
    statuses = Counter(row["status"] for row in records)
    return {
        "검증 통과 행": len(records),
        "입력 오류 행": len(issues),
        **{status: statuses[status] for status in STATUSES},
        "high 미흡": sum(row["status"] == "fail" and row["severity"] == "high" for row in records),
    }


def append_row(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            # 신뢰하지 않는 문자열은 수식 접두사와 무관하게 텍스트로 저장한다.
            cell.data_type = "s"


def build_workbook(records, issues, source_hash):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    summary = workbook.active
    summary.title = "요약"
    detail = workbook.create_sheet("점검 결과")
    errors = workbook.create_sheet("입력 오류")
    append_row(summary, ["보안 점검 결과 보고서", "값"])
    for label, value in summarize(records, issues).items():
        append_row(summary, [label, value])
    append_row(summary, ["원본 CSV SHA-256", source_hash])
    append_row(summary, ["집계 기준", "검증 통과 행만 집계. pass=양호, fail=미흡, review=검토, na=해당 없음"])
    append_row(summary, ["갱신 방법", "생성 시점의 집계입니다. CSV 수정 후 프로그램을 다시 실행하세요."])
    append_row(detail, [*COLUMNS, "원본 레코드 번호"])
    for row in records:
        append_row(detail, [*(row[key] for key in COLUMNS), row["source_record"]])
        detail.cell(detail.max_row, 1).number_format = "yyyy-mm-dd"
    append_row(errors, ["원본 레코드 번호", "오류 사유", "원본 값(JSON, 최대 30000자)"])
    for number, reason, raw in issues:
        append_row(errors, [number, reason, raw if len(raw) <= 30000 else raw[:29970] + "... [원본 CSV 참조]"])
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 22
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 85
    detail.column_dimensions["F"].width = 65
    errors.column_dimensions["B"].width = 50
    errors.column_dimensions["C"].width = 90
    return workbook


def verify_workbook(path, records, issues):
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    try:
        if workbook.sheetnames != ["요약", "점검 결과", "입력 오류"]:
            raise ValueError("시트 구조 오류")
        for name, count in (("점검 결과", len(records)), ("입력 오류", len(issues))):
            if workbook[name].max_row != count + 1:
                raise ValueError("행 수 불일치")
        for index, (label, count) in enumerate(summarize(records, issues).items(), 2):
            if workbook["요약"].cell(index, 1).value != label or workbook["요약"].cell(index, 2).value != count:
                raise ValueError("집계 불일치")
        if any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row):
            raise ValueError("허용하지 않은 수식 셀")
    finally:
        workbook.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args(argv)
    try:
        output = args.output.resolve()
        if output.suffix.lower() != ".xlsx" or output == args.input.resolve():
            raise ValueError("원본과 다른 .xlsx 출력 경로를 지정하세요.")
        records, issues, source_hash = load_checks(args.input)
        workbook = build_workbook(records, issues, source_hash)
        output.parent.mkdir(parents=True, exist_ok=True)
        # 기존 결과를 덮어쓰지 않는다. 재실행할 때 새 출력 이름을 사용한다.
        with output.open("xb") as stream:
            workbook.save(stream)
        workbook.close()
        verify_workbook(output, records, issues)
    except (OSError, ValueError, csv.Error, ImportError) as exc:
        print(f"보고서 생성 실패: {exc}")
        return 1
    print(f"보고서 생성 완료: {output}")
    print(f"검증 통과 {len(records)}건, 입력 오류 {len(issues)}건")
    return 2 if issues else 0


if __name__ == "__main__":
    raise SystemExit(main())
